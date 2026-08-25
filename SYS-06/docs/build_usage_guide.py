from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SYS-06_V2_Usage_Guide.docx"
BLUE = "2E74B5"
DARK = "1F4D78"
FILL = "E8EEF5"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    return p


def add_table(doc, rows, widths=(1.75, 4.75)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        shade(cells[0], FILL)
        set_cell_text(cells[0], label, bold=True, color=DARK)
        set_cell_text(cells[1], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.49)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7), ("Heading 3", 12, DARK, 10, 5)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("SYS-06 Operational Alerting - v2 Usage Guide")
    run.font.name = "Calibri"; run.font.size = Pt(9); run.font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("SYS-06 v2 | Internal operational reference")
    fr.font.name = "Calibri"; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor(100, 100, 100)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("SYS-06 v2 Usage Guide")
    r.font.name = "Calibri"; r.font.size = Pt(25); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    subtitle = doc.add_paragraph("Secure operational alert ingestion, delivery, monitoring, and testing")
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.runs[0].font.size = Pt(12); subtitle.runs[0].font.color.rgb = RGBColor(90, 90, 90)

    heading(doc, "1. What SYS-06 v2 does")
    doc.add_paragraph("SYS-06 v2 receives standardized alert payloads and delivers them to configured Telegram and Slack channels. It protects operators from alert fatigue through a five-minute duplicate window and a limit of ten outbound notifications per minute.")
    add_table(doc, [("HTTP endpoint", "POST /alerts"), ("Health endpoint", "GET /health"), ("Authentication", "X-Alert-Token is required when ALERT_DRY_RUN=false"), ("Rate limit", "10 outbound notifications per 60 seconds"), ("Duplicate window", "300 seconds per deduplicate_key")])

    heading(doc, "2. Before you start")
    bullet(doc, "Use Python 3.11+ or Docker Compose.")
    bullet(doc, "Keep real credentials only in SYS-06/.env. The .env.example file must remain credential-free.")
    bullet(doc, "Set ALERT_DRY_RUN=true first. This logs messages locally without sending Telegram or Slack notifications.")
    bullet(doc, "If a bot token was ever committed, revoke and regenerate it before using live mode.")

    heading(doc, "3. Configure the service")
    doc.add_paragraph("Copy the example configuration to .env and complete only the values needed for the chosen mode.")
    code(doc, "Copy-Item .env.example .env")
    add_table(doc, [("ALERT_DRY_RUN", "true for safe local testing; false for live delivery."), ("TELEGRAM_BOT_TOKEN", "BotFather token. Required for live Telegram delivery."), ("TELEGRAM_CHAT_ID", "Destination Telegram group or channel ID."), ("SLACK_WEBHOOK_URL", "Incoming webhook for live Slack delivery."), ("ALERT_API_TOKEN", "Long random secret required by POST /alerts in live mode."), ("LOG_LEVEL", "Use INFO normally; use DEBUG only while diagnosing issues.")])

    heading(doc, "4. Run with Docker")
    doc.add_paragraph("The v2 deployment is isolated from the original demo and uses Dockerfile.v2.")
    code(doc, "docker compose -f docker-compose.v2.yml up --build")
    doc.add_paragraph("Confirm that the service is healthy:")
    code(doc, "Invoke-RestMethod http://localhost:8080/health")
    doc.add_paragraph("Expected response fields include status, dry_run, queue_depth, and configured_channels.")

    heading(doc, "5. Send an alert")
    doc.add_paragraph("Use a JSON payload that follows the SYS-06 alert contract. In live mode, pass the API token in the X-Alert-Token header.")
    code(doc, "Invoke-RestMethod -Method Post -Uri http://localhost:8080/alerts -Headers @{ 'X-Alert-Token' = 'YOUR_SECRET' } -ContentType 'application/json' -Body '{\"alert_type\":\"INFRASTRUCTURE_FAILURE\",\"level\":\"CRITICAL\",\"source_system\":\"SYS-03\",\"event\":\"CONTAINER_DOWN\",\"deduplicate_key\":\"infra:kafka:down\",\"message\":\"Kafka container is unavailable\",\"channels\":[\"telegram\",\"slack\"]}'")
    heading(doc, "Payload fields", level=2)
    add_table(doc, [("alert_type", "One of the defined alert categories, for example INFRASTRUCTURE_FAILURE."), ("level", "CRITICAL, WARNING, or INFO."), ("source_system", "SYS-01 through SYS-06."), ("event", "Short UPPER_SNAKE_CASE event name."), ("deduplicate_key", "Stable incident key; use the same key for repeats of one incident."), ("message", "Operator-readable description."), ("channels", "Requested delivery channels, typically telegram and slack.")])

    heading(doc, "6. How routing and protection work")
    bullet(doc, "CRITICAL alerts target every configured real-time channel, so both Telegram and Slack normally receive them.")
    bullet(doc, "Each delivery counts against the limit. A critical alert sent to two channels uses two of the ten available slots.")
    bullet(doc, "When capacity is unavailable, the service queues the alert in FIFO order and retries it in the background.")
    bullet(doc, "An alert with a duplicate deduplicate_key is suppressed for five minutes, including while it is waiting in the queue.")
    bullet(doc, "A failure on one channel does not prevent the other configured channel from receiving the alert.")

    heading(doc, "7. Test before live deployment")
    code(doc, "E:\\Graduation_Thesis\\Test\\venv\\Scripts\\python.exe -m pytest E:\\Graduation_Thesis\\SYS-06\\tests -q")
    doc.add_paragraph("The expected result for the current v2 upgrade is: 14 passed. Run this after changes to alert rules, adapters, or payload models.")

    heading(doc, "8. Troubleshooting")
    add_table(doc, [("401 Invalid alert API token", "Set ALERT_API_TOKEN in .env and send the same value in X-Alert-Token. This check applies only in live mode."), ("No Telegram or Slack message", "Confirm ALERT_DRY_RUN=false, validate credentials, then inspect /health for configured_channels."), ("Alert is not sent again", "Check deduplicate_key. Repeated keys are intentionally suppressed for 300 seconds."), ("Alert is delayed", "Check queue_depth in /health. The ten-notification-per-minute limit has queued it."), ("Service will not start", "Run docker compose logs -f sys06-alert-dispatcher-v2 and check .env syntax.")])

    heading(doc, "9. Live-operation checklist")
    for item in [".env is present locally and is not committed.", "ALERT_DRY_RUN is false only after dry-run checks succeed.", "Telegram and Slack credentials are configured and tested.", "ALERT_API_TOKEN is long, unique, and supplied by alert producers.", "GET /health reports status ok and the expected channels.", "The full pytest suite passes before deployment."]:
        bullet(doc, item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
